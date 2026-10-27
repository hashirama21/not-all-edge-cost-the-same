"""Texte des documents (PDF, DOCX, HTML). Les PDF passent par PyMuPDF : texte, OCR des
pages scannées, tableaux et métadonnées."""

from __future__ import annotations

import logging
import shutil
import subprocess
import tempfile
from pathlib import Path

from selectolax.lexbor import LexborHTMLParser as HTMLParser

log = logging.getLogger(__name__)

MIN_CHARS_PER_PAGE = 80  # en dessous, la page est probablement une image
OCR_DPI = 300
OCR_LANG = "fra+eng"
MAX_TABLES = 200
IMAGE_SUFFIXES = (".jpg", ".jpeg", ".png", ".tif", ".tiff")


def _tessdata() -> str | None:
    """Dossier tessdata pour l'OCR intégré de PyMuPDF, ou None si Tesseract est absent."""
    import pymupdf

    try:
        return pymupdf.get_tessdata()
    except RuntimeError:
        return None


def _cell_text(words: list, bbox) -> str:
    """Mots dont le centre tombe dans la cellule, dans l'ordre de lecture.

    On recompose à partir des mots de PyMuPDF plutôt que de `Table.extract()`, qui colle
    les mots quand l'espace est étroit (« ETSALPHASARL »)."""
    if bbox is None:
        return ""
    x0, y0, x1, y1 = bbox
    inside = [w for w in words if x0 <= (w[0] + w[2]) / 2 <= x1 and y0 <= (w[1] + w[3]) / 2 <= y1]
    inside.sort(key=lambda w: (w[5], w[6], w[7]))
    return " ".join(w[4] for w in inside).strip()


def _to_markdown(rows: list[list[str]]) -> str:
    esc = lambda c: c.replace("|", "\\|").replace("\n", " ")
    lines = ["| " + " | ".join(esc(c) for c in rows[0]) + " |", "|" + "---|" * len(rows[0])]
    lines += ["| " + " | ".join(esc(c) for c in r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def _page_tables(page) -> list[dict]:
    """Tableaux détectés par PyMuPDF (lignes tracées, puis alignement du texte en repli)."""
    out = []
    words = page.get_text("words", sort=True)
    for strategy in ("lines", "text"):
        try:
            found = page.find_tables(vertical_strategy=strategy, horizontal_strategy=strategy)
        except Exception:
            continue
        for tab in found.tables:
            rows = [[_cell_text(words, c) for c in row.cells] for row in tab.rows]
            if tab.header.external:
                rows.insert(0, [_cell_text(words, c) for c in tab.header.cells])
            rows = [r for r in rows if any(r)]
            if len(rows) >= 2 and tab.col_count >= 2:
                out.append({"page": page.number + 1, "strategy": strategy, "header": rows[0],
                            "rows": rows[1:], "markdown": _to_markdown(rows)})
        if out:
            break
    return out


def pdf_text(path: Path, ocr: bool = True) -> tuple[str, dict]:
    """Texte page par page avec PyMuPDF ; les pages scannées passent par l'OCR intégré
    (Tesseract via PyMuPDF), sinon par ocrmypdf. Renvoie aussi tableaux et métadonnées."""
    import pymupdf

    meta = {"pages": 0, "ocr": False, "ocr_pages": 0, "scanned_pages": 0, "tables": [], "pdf_metadata": {}}
    try:
        doc = pymupdf.open(path)
    except Exception as exc:  # PDF corrompu
        log.warning("PDF illisible %s : %s", path, exc)
        return "", {**meta, "error": str(exc)}
    if doc.needs_pass and not doc.authenticate(""):
        doc.close()
        return "", {**meta, "error": "chiffré"}
    meta["pages"] = doc.page_count
    meta["pdf_metadata"] = {k: v for k, v in (doc.metadata or {}).items() if v}
    tessdata = _tessdata() if ocr else None
    pages = []
    for page in doc:
        text = page.get_text("text", sort=True)
        if len(text.strip()) < MIN_CHARS_PER_PAGE:
            meta["scanned_pages"] += 1
            if tessdata:
                try:
                    tp = page.get_textpage_ocr(language=OCR_LANG, dpi=OCR_DPI, full=True, tessdata=tessdata)
                    text = page.get_text("text", textpage=tp, sort=True)
                    meta["ocr_pages"] += 1
                except Exception as exc:
                    log.warning("OCR PyMuPDF en échec p.%s de %s : %s", page.number + 1, path.name, exc)
        if len(meta["tables"]) < MAX_TABLES:
            meta["tables"].extend(_page_tables(page))
        pages.append(text)
    doc.close()
    meta["ocr"] = meta["ocr_pages"] > 0
    if ocr and not tessdata and pages and meta["scanned_pages"] / len(pages) > 0.5:
        text = _ocr(path)
        if text:
            meta["ocr"] = True
            return text, meta
    return "\n\f".join(pages), meta


def pdf_markdown(text: str, meta: dict) -> str:
    """Entrée du LLM : texte suivi des tableaux PyMuPDF en Markdown (colonnes préservées)."""
    tables = meta.get("tables") or []
    if not tables:
        return text
    parts = [text, "\n\n# Tableaux détectés"]
    parts += [f"\n## Page {t['page']}\n{t['markdown']}" for t in tables]
    return "\n".join(parts)


def pdf_page_images(path: Path, dpi: int = 150, max_pages: int = 20) -> list[bytes]:
    """Pages rendues en PNG, pour envoyer un scan trop lourd au LLM page par page."""
    import pymupdf

    with pymupdf.open(path) as doc:
        return [page.get_pixmap(dpi=dpi).tobytes("png") for page in doc.pages(0, min(max_pages, doc.page_count))]


def _ocr(path: Path) -> str:
    """Repli : OCR via ocrmypdf (tesseract fra+eng), puis relecture avec PyMuPDF."""
    if not shutil.which("ocrmypdf"):
        log.info("ni Tesseract ni ocrmypdf : %s reste sans OCR", path.name)
        return ""
    import pymupdf

    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "ocr.pdf"
        cmd = ["ocrmypdf", "-l", OCR_LANG, "--skip-text", "--quiet", str(path), str(out)]
        if subprocess.run(cmd, capture_output=True).returncode != 0 or not out.exists():
            return ""
        with pymupdf.open(out) as doc:
            return "\n\f".join(page.get_text("text", sort=True) for page in doc)


def doc_text(path: Path) -> tuple[str, dict]:
    """Word 97-2003 (.doc) : antiword, sinon LibreOffice en mode sans interface."""
    if shutil.which("antiword"):
        res = subprocess.run(["antiword", "-m", "UTF-8.txt", str(path)], capture_output=True)
        if res.returncode == 0:
            return res.stdout.decode("utf-8", errors="replace"), {"converter": "antiword"}
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office:
        with tempfile.TemporaryDirectory() as tmp:
            res = subprocess.run([office, "--headless", "--convert-to", "txt:Text", "--outdir", tmp, str(path)],
                                 capture_output=True, timeout=300)
            out = Path(tmp) / (path.stem + ".txt")
            if res.returncode == 0 and out.exists():
                return out.read_text(encoding="utf-8", errors="replace"), {"converter": "libreoffice"}
    log.warning("%s : .doc non lu (installer antiword ou LibreOffice)", path.name)
    return "", {"unsupported": ".doc"}


def spreadsheet_text(path: Path) -> tuple[str, dict]:
    """Classeur Excel : une ligne de texte par ligne de feuille (« a | b | c »)."""
    import pandas as pd

    try:
        sheets = pd.read_excel(path, dtype=str, sheet_name=None, header=None)
    except Exception as exc:  # classeur corrompu ou format inattendu
        log.warning("classeur illisible %s : %s", path.name, exc)
        return "", {"error": str(exc)}
    parts = []
    for name, df in sheets.items():
        parts.append(f"# {name}")
        parts += [" | ".join(v for v in row if v) for row in df.fillna("").to_numpy().tolist() if any(row)]
    return "\n".join(parts), {"sheets": list(sheets)}


def image_text(path: Path) -> tuple[str, dict]:
    """Image (communiqué scanné) : OCR par PyMuPDF + Tesseract si disponible."""
    import pymupdf

    tessdata = _tessdata()
    if not tessdata:
        log.info("Tesseract absent : %s reste sans OCR", path.name)
        return "", {"ocr": False, "unsupported": path.suffix.lower()}
    try:
        with pymupdf.open(path) as img, pymupdf.open("pdf", img.convert_to_pdf()) as doc:
            page = doc[0]
            tp = page.get_textpage_ocr(language=OCR_LANG, dpi=OCR_DPI, full=True, tessdata=tessdata)
            return page.get_text("text", textpage=tp, sort=True), {"ocr": True, "ocr_pages": 1, "pages": 1}
    except Exception as exc:
        log.warning("OCR impossible sur %s : %s", path.name, exc)
        return "", {"ocr": False, "error": str(exc)}


def docx_text(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def html_text(html: str) -> str:
    tree = HTMLParser(html)
    for node in tree.css("script, style, noscript, nav, footer, header"):
        node.decompose()
    body = tree.body or tree.root
    return body.text(separator="\n", strip=True) if body else ""


def html_tables(html: str) -> list[list[dict]]:
    """Chaque <table> devient une liste de dicts {en-tête: cellule}."""
    tree = HTMLParser(html)
    out = []
    for table in tree.css("table"):
        rows = table.css("tr")
        if not rows:
            continue
        headers = [c.text(strip=True) for c in rows[0].css("th, td")]
        records = []
        for row in rows[1:]:
            cells = [c.text(strip=True) for c in row.css("td, th")]
            if cells and any(cells):
                keys = headers if len(headers) == len(cells) else [f"col{i}" for i in range(len(cells))]
                records.append(dict(zip(keys, cells)))
        out.append(records)
    return out


def document_text(path: Path, content_type: str = "", ocr: bool = True) -> tuple[str, dict]:
    suffix = path.suffix.lower()
    if suffix == ".pdf" or "pdf" in content_type:
        return pdf_text(path, ocr=ocr)
    if suffix == ".docx":
        return docx_text(path), {}
    if suffix == ".doc" or content_type.startswith("application/msword"):
        return doc_text(path)
    if suffix in (".xlsx", ".xls"):
        return spreadsheet_text(path)
    if suffix in IMAGE_SUFFIXES:
        return image_text(path)
    if suffix in (".html", ".htm") or "html" in content_type:
        return html_text(path.read_text(encoding="utf-8", errors="replace")), {}
    if suffix in (".txt", ".json"):
        return path.read_text(encoding="utf-8", errors="replace"), {}
    return "", {"unsupported": suffix}
